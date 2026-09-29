"""Build the evidence-based Chinese report with python-docx in an artifact venv."""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT/'logs/bench/v52_sim2sim_root_cause'
OUTPUT = ROOT/'docs/push_cart_v52_sim2sim_resolution.docx'
DATA = json.loads((EVIDENCE/'final_summary.json').read_text(encoding='utf-8'))
doc = Document()


def font(style, size, bold=False):
    s = doc.styles[style]
    s.font.name = 'Microsoft YaHei'
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.italic = False
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.line_spacing = 1.07


for s, size, bold in [('Normal', 10.5, False), ('Title', 22, True),
                      ('Subtitle', 10.5, False), ('Heading 1', 16, True),
                      ('Heading 2', 12, True), ('Caption', 9, False)]:
    font(s, size, bold)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.styles['Normal'].paragraph_format.widow_control = True
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2)
sec.footer_distance = Cm(.8)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run('V5.2 仿真验证  ·  ')
r.font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'V5.2 推车估力的 sim2sim 误差排查与修正报告'
doc.core_properties.subject = '双手抓握 动力学参数 力标签 控制接口 训练数据和固定网络验证'
doc.core_properties.author = 'Codex'


def p(text, style=None, bold=False):
    x = doc.add_paragraph(style=style)
    x.add_run(text).bold = bold
    return x


def h(text, level=1):
    doc.add_heading(text, level)


def page(title):
    paragraph = doc.add_heading(title, 1)
    paragraph.paragraph_format.page_break_before = True


def table(headers, rows, widths, size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement('w:'+name)
        for k, v in {'val':'single', 'sz':'4', 'color':'D9D9D9'}.items(): b.set(qn('w:'+k), v)
        borders.append(b)
    pr.append(borders)
    for col, width in zip(t.columns, widths): col.width = Cm(width)
    for i, row in enumerate([headers]+rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, (cell, text) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(text)
            tcpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'DCE6F1' if i == 0 else ('F5F7F9' if i % 2 == 0 else 'FFFFFF'))
            tcpr.append(shd)
            margins = OxmlElement('w:tcMar')
            for name, val in [('top','70'), ('bottom','70'), ('left','95'), ('right','95')]:
                x = OxmlElement('w:'+name); x.set(qn('w:w'), val); x.set(qn('w:type'), 'dxa'); margins.append(x)
            tcpr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.08
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.size = Pt(size); r.bold = i == 0
                    r.font.color.rgb = RGBColor(0,0,0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.add_run().font.size = Pt(2)
    return t


def picture(path, caption, width=6.45):
    x = doc.add_paragraph()
    x.paragraph_format.keep_with_next = True
    x.add_run().add_picture(str(path), width=Inches(width))
    for inline in x._p.xpath('.//wp:docPr'):
        inline.set('descr', caption)
    p(caption, 'Caption')


def code(text):
    x = p(text)
    x.paragraph_format.space_after = Pt(7)
    x.paragraph_format.line_spacing = 1.0
    for r in x.runs:
        r.font.name = 'Consolas'; r.font.size = Pt(8.3)
    return x


def values(head, condition):
    return DATA['fixed_heads'][head][condition]['handle']['rmse']


def triple(v):
    return ' / '.join(f'{x:.3f}' for x in v)


p('V5 2 推车估力的 sim2sim\n误差排查与修正报告', 'Title')
p('实验与代码记录  2026年9月27日', 'Subtitle')
p('本报告解释 V5.2 的力估计器为什么在 Isaac 中较准、迁移到 MuJoCo 后误差增大，并记录已经落实的抓握和惯量修正。验证始终围绕完整的仿真推车：机器人、双手固定抓握、车体、车轮、地面接触和闭环控制均保留。')
p('目前最强的因果证据指向双手抓握柔顺性。原 MuJoCo weld 与 Isaac 固定关节产生了明显不同的左右手内力；在不改变估力网络的情况下，仅调整抓握就显著降低了预测误差。另一个已确认的问题是实际 USD／PhysX 惯量与先前使用的 URDF 惯量不一致，现已按实际训练场景兼容。', bold=True)
p('但跨引擎误差仍未消除。网络输入、短历史、训练覆盖和驱动数值也有影响；没有证据支持“估力器完全没有问题”或“已完整复现 FALCON”。以下数值分别属于不同网络与驱动条件，不能混合比较。')
table(['固定网络与驱动', 'Isaac RMSE', '最终 MuJoCo RMSE'], [
    ['原始 checkpoint\n混合驱动', triple(values('original','isaac')), triple(values('original','final'))],
    ['重训全输入 680 维\n显式 PD', triple(values('full','isaac')), triple(values('full','final'))],
    ['重训全身历史 435 维\n显式 PD', triple(values('fullbody','isaac')), triple(values('fullbody','final'))],
], [5.2,5.7,6.1])
p('表中依次为把手合成 Fx、Fy、Mz 的 RMSE，单位 N、N、N·m。每组 12 个工况、4,800 个评分样本，均无工况失败。原始 checkpoint 的默认入口不会自动换成重训力头。', 'Caption')
h('完成范围', 2)
p('已完成原始方案对照、标签与采样核验、控制和观测核查、抓握干预、运行时惯量核验、三组固定网络复测，以及 33 项相关回归测试。确定的修正已接入 V5.2 对齐入口；历史配置和失败诊断记录保留。第八部分单列仍然存在的误差与验证边界。')

page('一  双手闭链为什么放大抓握差异')
p('两只手同时固定在一个把手上时，手臂和车构成闭链。两手可以互相拉扯，使左手与右手受力都很大，而车受到的合力仍然很小。因此，车看起来运动正常，并不能证明左右手负载分配已经一致。')
p('在移除学习策略、重力、地面接触和刚体阻尼的隔离实验中，两引擎执行相同关节目标和已知车体外载荷。目标最大差仅 8.00×10⁻⁸ rad。单手连接时差异较小，双手连接时逐手差异跃升，合力差却远小于逐手差。')
table(['无策略力向量差 RMS', '原 MuJoCo weld', '仅修正抓握'], [
    ['左手单独抓握', '0.753 N', '0.161 N'],
    ['右手单独抓握', '0.788 N', '0.197 N'],
    ['双手抓握 左手／右手', '23.852 / 23.904 N', '2.758 / 2.766 N'],
    ['双手抓握 合力', '1.293 N', '0.284 N'],
], [7.1,4.9,5.0])
picture(EVIDENCE/'dual_grasp_response.png', '图 1  相同目标和外载荷下的受力响应。左右手侧向力可以相互抵消；本图是物理响应对照，不是网络预测误差。', width=6.05)
p('处理方式：只将两条 grasp equality 的 solimp 前两项由 0.95／0.99 改为 MuJoCo 文档允许的上限 0.9999，形成 rigid_limit 配置。时间步、solref、PD 和接触参数不随这项干预改变。参数先由独立物理探针验证，再评测网络，未按预测分数搜索。')

page('二  惯量差异的来源与兼容处理')
h('机器人实际使用的惯量来自组合后的 USD', 2)
p('机器人关节几何与 URDF 基本一致，但部分 USD 刚体的主惯量方向不同。以左膝为例，运行时与 URDF 张量的矩阵范数差约 0.00909 kg·m²。仅核对总质量、关节位置或惯量特征值，都会漏掉方向差异。')
p('隔离各向异性刚体，施加三个方向的已知力矩并测角加速度，反求出的物理惯量与 PhysX get_inertias 的矩阵差约 4.01×10⁻⁷ kg·m²。组合 USD 重建也与 API 一致。这排除了“只是读错坐标系”的解释。')
h('车载荷重置覆盖了更新后的主轴', 2)
p('旧训练代码先设置惯量，再恢复原 COM pose。COM pose 不仅包含位置，也包含惯量主轴的旋转；后一次写入覆盖了 set_inertias 刚计算出的主轴。这个顺序作用于全部车刚体，包括车轮。')
table(['额外载荷', '原意图的车体惯量对角', '实际 PhysX 惯量对角'], [
    ['0 kg', '0.6516 / 1.3600 / 1.2916', '0.6516 / 1.2916 / 1.3600'],
    ['12.5 kg', '0.7870 / 1.5683 / 1.5520', '0.7870 / 1.5520 / 1.5683'],
    ['25 kg', '0.9224 / 1.7767 / 1.8124', '0.9224 / 1.8124 / 1.7767'],
], [2.3,7.3,7.4])
p('单位 kg·m²；该表仅展示居中载荷。偏心载荷必须保留实际主轴处理，不能简单套用一个通用的 y／z 交换。', 'Caption')
h('这次对齐已有训练域', 2)
p('已导出组合 USD 的 49 个机器人刚体和 9 个车刚体，保存质量、COM、主惯量、主轴和来源 SHA256。MuJoCo 按验证过的固定关节变换合并刚体，并在载荷更新时复现旧 PhysX 的主轴覆盖行为；车轮也一并处理。运行入口无需安装 USD 解析包，源资产变化时会拒绝使用过期参考。')
p('直接调用实际 Isaac 重置函数，测试 0、5、12.5、25、50 kg 以及居中／平面偏心，共 20 组，最大张量差 1.32×10⁻⁶ kg·m²。标准三载荷下的全部刚体对照最大差为 1.10×10⁻⁶ kg·m²。')
p('这项修正使物理参数更忠实于旧训练场景，但没有让所有预测指标继续下降。因此它是已证实并修正的参数问题，尚不能称为主要预测误差来源。未来若修正 Isaac setter 顺序，应单独版本化物理场景并重采、重训，不能默认为旧 checkpoint 已适应该变化。')

page('三  标签代表什么以及 FALCON 的差别')
p('目标是估计推车时左右手各自施加在车上的三维力和三维力矩，共 12 维。它与车的加速、转弯、轮阻和碰撞有关，也包括双手之间的内力。把手合成的 Fx、Fy、Mz 是由逐手预测按固定抓握几何合成的任务指标。')
table(['项目', '本项目推车估力', 'FALCON 附录 A3 对照'], [
    ['监督对象', '双手 12 维力与力矩', '三维末端外力'],
    ['标签来源', '物理交互反力重建', '已知施加外力'],
    ['关键输入', '对照包含全输入及全身历史', '全身 q／dq、角速度、重力\n及过去下肢动作'],
    ['训练关系', '冻结行为策略后单独训练力头', '估计器与下身策略联合训练'],
    ['使用预测', '接回已有高层 Actor', '还包含 JᵀF 上身补偿'],
], [2.5,7.4,7.1], size=9.2)
p('FALCON 的公开 MuJoCo 分支将已知基座系力旋转到世界系，再写入手部 xfrc_applied。训练端也保存实际施加力进行比较。这保证了该已知外力试验的物理含义一致，但不能直接替代双手推车闭链的反力标签。')
h('两引擎采用统一物理定义', 2)
p('Isaac 从手部子树的 incoming force 中扣除惯性、重力等项重建负载；MuJoCo 从双手 weld 约束作用重建负载。两端统一为“手对车”，力矩参考点为各手掌原点，坐标轴为车的航向坐标系。每个 20 ms 控制步，评分标签取四个 5 ms 子步的平均。标签只用于训练监督和评分，不作为估力器的真值输入。')
h('基础标签检查已扩展到实际推车', 2)
p('相同已知小外力实验中，Isaac 最大组力向量 RMS 约 0.002714 N，MuJoCo 独立 RNE 核验约 6.8×10⁻¹⁴ N。这些是标签重建残差，不是网络精度，也不是完整推车场景的充分证明。')
p('实际闭环推车另检查 12 工况、19,200 个物理时刻。手掌标签与各自肩部力减去整条手臂惯性／重力的独立通道重建，力分量 RMS 为 0.00071–0.00179 N，最大瞬时差 0.00705 N。它支持逐手力标签内部一致，不能等同独立实体力传感器标定。')
p('车侧含地面接触的整体平衡，力残差 RMS 为 0.031／0.034／0.103 N，力矩残差为 0.275／0.453／0.035 N·m。力矩与瞬态仍有残差，报告没有把全部 12 维标签称为无误差真值。')

page('四  网络和训练数据从哪里来')
p('独立力头采用本项目的 ELU 多层感知机，隐藏层为 512、256、128，输出双手 12 维归一化 wrench，使用独立 Huber 监督。FALCON 提供了输入与训练思路的参照；本网络不是从其官方完整附录训练实现中直接复制的模型。')
table(['输入版本', '最近五帧内容', '总维数'], [
    ['全输入', '关节／IMU、力矩代理、DLS 提示、\nHIM、固定抓握 FK、命令等', '680'],
    ['A3 输入对照', 'q 27＋dq 27＋角速度 3＋重力 3\n＋过去下肢动作 12', '360'],
    ['全身控制历史', 'A3 输入对照再加入每帧\n实际执行的上身目标 15', '435'],
], [3.3,11.5,2.2])
p('全身历史版本已经包含上身控制目标。三种版本的 q／dq 和 IMU 都是全身本体观测，并非只观察下肢。全输入中的阻尼最小二乘结果是物理提示，既不是监督标签，也不能把“有 DLS”直接解释成误差原因。')
h('样本来自 Isaac 中实际仿真推车', 2)
p('使用冻结 HOMIE 和既定高层行为策略，在显式 PD 的 Isaac 推车环境中采集。每一步先由旧观测产生动作，完成四个物理子步，再把新观测与刚完成区间的受力标签一起保存。主数据为 64 环境 × 4096 控制步，剔除无效／重置过渡后保留 246,208 条。')
table(['划分', '样本数', '划分方式'], [
    ['训练', '184,656', '完整环境划分'], ['验证', '30,776', '完整环境划分'],
    ['Isaac 测试', '30,776', '完整环境划分'],
], [4.0,4.0,9.0])
p('MuJoCo 未参与主模型的归一化、早停或模型选择。11／22／33 是三个优化种子，共用一次采集数据；它们不是三个独立采集种子。报告中最终物理干预保持选定权重不变，没有再训练来掩盖场景差异。')
h('采集核查结论', 2)
p('245,888 对同回合同环境的相邻样本，历史重叠逐项相同，未发现错位或环境／回合／age 重复。两端在相同传感样本上的实际观测构造也通过一致性检查。全身历史中的上身目标已与实际提交的 PD 目标核对。')
p('旧主数据缺少逐样本实际载荷等审计字段，不能追溯宣称所有载荷覆盖完整。采集器已补充车总质量、命令、有效子步数和 27 关节目标；新增 1,656 条短采样验证了字段正确性，但不替代多种子、全工况覆盖验证。')

page('五  固定网络修改前后的完整比较')
p('每组为 0／12.5／25 kg 额外载荷，分别执行 0.3／0.6 m/s 直行及左右转弯，共 12 工况。每工况 500 个控制步，舍弃前 2 s，累计 4,800 个评分样本。空车总质量 14.4 kg；噪声、延迟和域随机化关闭。两引擎分别闭环运行，轨迹并非逐状态相同。')
rows = []
for head, label in [('original','原始混合驱动'), ('full','全输入显式 PD'), ('fullbody','全身历史显式 PD')]:
    for cond, cname in [('isaac','Isaac'), ('mujoco_nominal','旧 MJ 抓握'), ('grasp_only','仅改抓握'), ('final','最终 MJ')]:
        rows.append([label, cname]+[f'{v:.3f}' for v in values(head,cond)])
table(['固定网络', '场景', 'Fx RMSE\nN', 'Fy RMSE\nN', 'Mz RMSE\nN·m'], rows, [5.0,3.6,2.8,2.8,2.8], size=9.1)
p('旧 MJ 抓握行指此前已做资产／控制参数对齐的场景，不是更早的 legacy XML。仅改抓握和最终 MJ 之间只增加已验证的运行时惯量兼容。', 'Caption')
h('实际左右手预测精度', 2)
p('最终全输入力头的逐手分量 RMSE 如下。把手 Fx 的 4.788 N 与左手 Fx 的 3.195 N 是不同指标，不能互相替代。RMSE 也不是“每一时刻误差都小于这个数”。')
v = DATA['fixed_heads']['full']['final']['hands']['rmse']
table(['手', 'Fx N', 'Fy N', 'Fz N', 'Mx N·m', 'My N·m', 'Mz N·m'],
      [['左']+[f'{x:.3f}' for x in v[:6]], ['右']+[f'{x:.3f}' for x in v[6:]]],
      [2.0,2.5,2.5,2.5,2.5,2.5,2.5], size=9.2)
p('此前提到的 1.72，是重训全输入力头在 Isaac 中的把手 Fx RMSE。它说明该域、该组工况预测较好，不能当成两只手所有分量的统一误差，更不能作为 MuJoCo 上的成绩。')
p('原始 checkpoint 的原驱动域与显式 PD 重训域不同。既有对照中，在 Isaac 内直接把旧模型切成显式 PD，Fx RMSE 也会升至约 14.19 N；驱动方式改变本身就改变网络观测与动力学，不能只换部署 PD 而沿用原精度结论。')

page('六  实际预测曲线与剩余输入分布差异')
picture(EVIDENCE/'final_prediction_example.png', '图 2  最终全输入力头在 MuJoCo 的实际预测。示例为额外载荷 0 kg、0.3 m/s 前进工况，展示整个评分区间；未追加平滑或挑选局部最优片段。蓝线是约束参考，橙线是冻结网络预测。')
p('用保存的观测重新执行同一力头，重算 RMSE 与在线记录相符。曲线能反映步态周期受力变化，但仍存在幅值和局部瞬态偏差。这比单独报一个综合 RMSE 更能说明当前预测情况。')
h('抓握修正没有消除所有输入域差异', 2)
rows = []
for name, label in [('imu_gyro','IMU 角速度'), ('joint_velocity','关节速度'),
                    ('joint_position','关节位置'), ('upper_target_applied','上身控制目标')]:
    rows.append([label]+[f"{DATA['coverage'][k][name]:.2f}%" for k in ('nominal','grasp_only','final')])
table(['最新帧字段', '旧抓握', '仅改抓握', '最终配置'], rows, [6.2,3.6,3.6,3.6])
p('数值为超过 Isaac 训练环境逐坐标 1%–99% 分位区间的标量比例，分母是样本数乘字段维数。使用同一全输入力头轨迹及其训练集；不是独立轨迹失败率，也不是误差因果占比。单坐标覆盖不保证多维组合覆盖。', 'Caption')
p('全输入力头最终 vx／vy／wz 跟踪 RMSE 为 0.096／0.093／0.102，单位 m/s、m/s、rad/s。估力改善与控制跟踪改善不是同一结论，尤其不能据此宣布所有转弯、载荷和扰动任务已解决。')

page('七  已落实的代码与复现入口')
p('正式评测链条为 eval_pushcart_v52_sim2sim.py → V5xPolicy／V5xRunner → 共享 MuJoCo 环境。V5.2 对齐入口默认 rigid_limit 抓握＋physx 惯量；V5.1、legacy 场景和基础诊断环境保持各自原配置。未修改训练资产、权重或历史结果。')
table(['代码位置', '处理内容'], [
    ['pushcart_bench/sim2sim_scene.py', '生成对齐场景、抓握参数、载荷与惯量更新'],
    ['pushcart_bench/physx_inertia.py', '验证 USD 来源并复现旧 PhysX 重置惯量'],
    ['pushcart_bench/policies_v5x.py', 'V5.2 默认选择与历史配置兼容'],
    ['scripts/eval_pushcart_v52_sim2sim.py', '显式选项、有效配置报告和独立结果命名'],
], [8.3,8.7], size=9.0)
p('惯量参考来自 pushcart_bench/data/v52_usd_mass_properties.json，由 export_v52_usd_mass_properties.py 从组合 USD 生成。源层与固定关节 URDF 均带哈希校验；新的结果文件名包含实际抓握和惯量配置，避免覆盖历史对照。')
h('运行默认原始 checkpoint', 2)
code('& E:\\envs\\RLisaac601\\python.exe `\n  scripts/eval_pushcart_v52_sim2sim.py `\n  --engine mujoco --profiles isaac --capture `\n  --output logs/bench/v52_reproduce_original')
h('运行已验证的重训全输入力头', 2)
code('$forceHead = "logs/bench/v52_falcon_explicit/" + `\n  "model_explicit_pd_full_22/estimator.pt"\n& E:\\envs\\RLisaac601\\python.exe `\n  scripts/eval_pushcart_v52_sim2sim.py `\n  --engine mujoco --profiles isaac --upper_drive explicit_pd `\n  --estimator $forceHead --audit_labels --capture `\n  --output logs/bench/v52_reproduce_full')
p('上面两个命令会得到不同网络的结果。第二个命令显式选择重训网络及其驱动域；不能省略这些参数后仍声称复现了 4.79 N 的成绩。')
h('保留对照与回归验证', 2)
p('在对应命令中加入 --inertia_profile urdf，只保留抓握修正；再加入 --grasp_profile nominal，复现旧对齐场景。历史跨版本比较继续使用 legacy。新增回归检查覆盖抓握参数可逆性、已测膝关节／偏心车体／车轮惯量，以及原有观测、标签、控制和兼容性；共 33 项通过。')

page('八  当前结论和验证边界')
h('已经有实验证据支持的判断', 2)
p('抓握柔顺性是已定位的重要误差来源：无网络物理探针复现了逐手负载分配差，固定网络的完整推车评测又验证了干预效果。实际 USD 惯量和旧重置顺序也与先前 MuJoCo 不同，现已对齐到旧模型真正经历的仿真域。')
p('标签的物理定义、符号、坐标和基础时序已有交叉核查。原始数据没有发现历史错位或未来标签泄漏的证据；新增采集字段已验证可用。全身历史模型已纳入上身控制，但在 Isaac 域内仍弱于全输入模型，不能排除输入信息量、可观测性和学习目标的影响。')
h('不能从本轮实验推出的结论', 2)
p('修正后的两引擎不是完全等价。PhysX TGS 与 MuJoCo 的约束／接触、驱动离散化和积分过程仍不同；无策略双手探针还有约 2.76 N 的逐手力向量差。原始 checkpoint 与重训网络也仍存在跨域性能缺口。')
p('实际推车平衡检查是内部一致性证据，力矩和部分瞬态仍有残差；它不是外部传感器认证。当前模型选择只使用一个采集种子和有限载荷／命令网格，也未证明极端载荷、松脱抓握、碰撞、长时间运行或真机性能。')
p('完整的 FALCON 联合训练与 JᵀF 补偿尚未复现。当前全输入力头也包含 DLS、力矩代理和任务信息，不能把其成绩表述为“仅凭纯本体历史的 FALCON 结果”。')
h('后续验证应如何承接', 2)
p('下一轮应固定当前物理基线，使用独立采集种子和完整工况留出，补齐载荷／命令／激励覆盖；比较更长的全身历史及状态辅助监督。已知施力的开链探针继续用于定位纯估力误差，完整推车用于验证双手内力和闭环效果，两者分别报告。')
p('如果修正 Isaac 的惯量 setter 顺序，则应建立新物理版本，先核验读回惯量和已知力矩响应，再采集并训练。已训练模型的兼容场景继续保留，避免把物理变化、网络变化和控制变化同时混入一个对比。')
p('本轮已完成确定问题的修正与复测，形成可复现的默认仿真入口；“所有 sim2sim 误差已被彻底消除”并不成立。验收依据是已知缺陷修复、物理核验通过和诚实保留残余误差，而不是强行将两个引擎的评分调到一致。', bold=True)

page('九  原始来源与证据索引')
p('项目基线为 8deff8a v5.2。本报告对应本次工作区修改和 2026年9月27日保存的实验结果；未创建提交。以下路径均相对仓库 E:\\CJME_Push_car。')
h('论文与官方实现', 2)
p('FALCON 附录 A.3\nhttps://arxiv.org/html/2505.06776v2#A1.SS3\n官方仓库 https://github.com/LeCAR-Lab/FALCON\n核对提交 a967a6d8494f57777cf8d266a644ac8e45833301')
p('MuJoCo 官方约束参数说明\nhttps://mujoco.readthedocs.io/en/stable/modeling.html#solver-parameters\n阻抗上限 0.9999；正 solref 时间常数默认不小于两倍物理步长。')
p('PhysX 主轴分解实现\nhttps://github.com/NVIDIA-Omniverse/PhysX\nphysx/source/foundation/FdMathUtils.cpp 中 PxDiagonalize；本项目另外用实际运行时读回和已知力矩响应验证，未仅依赖源码推断。')
p('其他参考：Learning Force Control for Legged Manipulation\nhttps://arxiv.org/html/2405.01402\nUniFP https://arxiv.org/html/2505.20829\n用于对照全身历史、辅助状态监督和激励设计；未宣称完整复现。')
h('本地报告与机器可读结果', 2)
for text in [
    '总根因记录：docs/push_cart_v52_falcon_label_root_cause.md',
    '最终结果与每组来源哈希：logs/bench/v52_sim2sim_root_cause/final_summary.json',
    '抓握独立探针：同目录 grasp_response_mujoco/、grasp_response_high_impedance/ 和 grasp_response_usd_inertia/',
    '真实推车标签：同目录 walking_dynamics/report.json 和 traces.npz',
    '运行时惯量：同目录 runtime_parameters_physx_inertia_final.json、cart_inertia_reset.json 和 inertia_frame.json',
    '训练与采集审计：docs/push_cart_v52_force_sources_and_data_audit.md；logs/bench/v52_falcon_explicit/falcon_a3/dataset_audit.json',
    '原始模型最终复测：同目录 original_checkpoint_final_alignment/；全输入：full_physx_inertia/；全身历史：fullbody_final_alignment/',
]:
    x=p(text); x.paragraph_format.space_after=Pt(7)
    for r in x.runs: r.font.size=Pt(9.5)
p('这里“同目录”均指 logs/bench/v52_sim2sim_root_cause。历史失败诊断未删除；有效性问题及修正理由见根因与采集审计文档。最终成绩来自完整成功的工况，而非筛掉失败后的子集。', 'Caption')

# Word's default Title style can retain a colored border even after font
# overrides. Remove border residue throughout the styles and document.
for container in (doc.styles.element, doc.element):
    for border in container.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(OUTPUT)
print(OUTPUT)
print('paragraphs', len(doc.paragraphs), 'tables', len(doc.tables), 'figures', len(doc.inline_shapes))
